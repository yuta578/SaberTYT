import os
import sys
import io

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import django
from decimal import Decimal

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'core.settings')
django.setup()

from django.test import Client
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from simulador.models import (
    Usuario, Categoria, Pregunta, OpcionRespuesta,
    Cuestionario, CuestionarioPregunta, Intento, RespuestaIntento, Importacion
)
import openpyxl

def audit_system():
    client = Client()
    errors = []
    successes = []

    def check(name, condition, err_msg=""):
        if condition:
            successes.append(name)
            print(f"  [OK] {name}")
        else:
            errors.append(f"{name}: {err_msg}")
            print(f"  [FAIL] {name}: {err_msg}")

    print("======================================================================")
    print("AUDITORIA INTEGRAL DE MODULOS FUNCIONALES - SABER TyT")
    print("======================================================================\n")

    # ------------------------------------------------------------------
    # 1. INTEGRIDAD DE LA BASE DE DATOS Y DATOS SEMILLA
    # ------------------------------------------------------------------
    print("1. VERIFICANDO INTEGRIDAD DE LA BASE DE DATOS Y MODELOS...")
    total_usuarios = Usuario.objects.count()
    check("BD: Usuarios existentes", total_usuarios >= 3, f"Total encontrados: {total_usuarios}")

    total_categorias = Categoria.objects.count()
    check("BD: Competencias / Categorias registradas", total_categorias >= 4, f"Total encontradas: {total_categorias}")

    total_preguntas = Pregunta.objects.count()
    check("BD: Banco de preguntas reales", total_preguntas >= 15, f"Total encontradas: {total_preguntas}")

    total_opciones = OpcionRespuesta.objects.count()
    check("BD: Opciones de respuesta", total_opciones >= 40, f"Total encontradas: {total_opciones}")

    total_cuestionarios = Cuestionario.objects.count()
    check("BD: Cuestionarios estructurados", total_cuestionarios >= 3, f"Total encontrados: {total_cuestionarios}")

    admin_user = Usuario.objects.filter(rol='admin').first()
    docente_user = Usuario.objects.filter(rol='docente').first()
    estudiante_user = Usuario.objects.filter(rol='estudiante').first()

    check("BD: Usuario Administrador disponible", admin_user is not None)
    check("BD: Usuario Docente disponible", docente_user is not None)
    check("BD: Usuario Estudiante disponible", estudiante_user is not None)

    # ------------------------------------------------------------------
    # 2. MODULO DE AUTENTICACION Y SEGURIDAD (M01, M03)
    # ------------------------------------------------------------------
    print("\n2. VERIFICANDO MODULO DE AUTENTICACION, REGISTRO Y PERFIL...")
    # Login GET
    resp = client.get(reverse('login'))
    check("Auth: Vista Login GET 200", resp.status_code == 200)

    # Login POST con credenciales validas
    resp = client.post(reverse('login'), {
        'email': admin_user.email,
        'password': 'demo1234'
    }, follow=True)
    check("Auth: Login Administrador redirige y crea sesion", resp.status_code == 200 and client.session.get('demo_profile', {}).get('role') == 'admin')

    # Logout
    resp = client.get(reverse('logout'), follow=True)
    check("Auth: Logout limpia la sesion", client.session.get('demo_profile') is None)

    # Registro de nuevo estudiante con generacion de codigo aleatorio
    nuevo_email = f"estudiante_test_{os.urandom(3).hex()}@itp.edu.co"
    resp = client.post(reverse('register'), {
        'nombre_completo': 'Estudiante Test Validacion',
        'email': nuevo_email,
        'password': 'password123',
        'confirm_password': 'password123',
        'telefono': '3001234567',
        'programa_academico': 'Tecnologia en Gestion Contable'
    }, follow=True)
    nuevo_creado = Usuario.objects.filter(email=nuevo_email).first()
    check("Auth: Registro crea usuario en MySQL", nuevo_creado is not None)
    if nuevo_creado:
        check("Auth: Codigo asignado aleatorio (EST-XXXXXXXX)", nuevo_creado.codigo_institucional.startswith('EST-') and len(nuevo_creado.codigo_institucional) == 12)
        check("Auth: Rol forzado a 'estudiante'", nuevo_creado.rol == 'estudiante')

    # Perfil GET y POST
    s = client.session
    s['demo_profile'] = {
        'id': nuevo_creado.id if nuevo_creado else estudiante_user.id,
        'role': 'estudiante',
        'name': 'Estudiante Test',
        'code': 'EST-12345678',
        'email': nuevo_email,
        'telefono': '3001234567',
        'programa': 'Tecnologia en Gestion Contable'
    }
    s.save()
    resp = client.get(reverse('profile'))
    check("Auth: Vista Mi Perfil GET 200", resp.status_code == 200)

    resp = client.post(reverse('profile'), {
        'telefono': '3109876543',
        'programa_academico': 'Tecnologia en Sistemas'
    }, follow=True)
    check("Auth: Actualizacion de datos de perfil", resp.status_code == 200)

    # ------------------------------------------------------------------
    # 3. MODULO ADMINISTRADOR (M02, M04, M05, M06, M12)
    # ------------------------------------------------------------------
    print("\n3. VERIFICANDO MODULO ADMINISTRADOR...")
    s = client.session
    s['demo_profile'] = {
        'id': admin_user.id,
        'role': 'admin',
        'name': admin_user.nombre_completo,
        'code': admin_user.codigo_institucional,
        'email': admin_user.email,
        'redirect': 'admin_dashboard'
    }
    s.save()

    # Dashboard Admin
    resp = client.get(reverse('admin_dashboard'))
    check("Admin: Dashboard General GET 200", resp.status_code == 200)

    # Directorio de usuarios
    resp = client.get(reverse('admin_users'))
    check("Admin: Directorio de Usuarios GET 200", resp.status_code == 200)

    # Crear usuario desde admin
    email_crear = f"docente_admin_{os.urandom(3).hex()}@itp.edu.co"
    resp = client.post(reverse('admin_crear_usuario'), {
        'name': 'Docente Creado Admin',
        'email': email_crear,
        'role': 'docente',
        'password': 'password123',
        'telefono': '3201112233',
        'programa_academico': 'Ciencias Basicas'
    }, follow=True)
    doc_creado = Usuario.objects.filter(email=email_crear).first()
    check("Admin: Creacion de usuario por administrador", doc_creado is not None and doc_creado.codigo_institucional.startswith('DOC-'))

    # Matriz de permisos
    resp = client.get(reverse('admin_permissions'))
    check("Admin: Matriz de Permisos GET 200", resp.status_code == 200)

    # Configuracion global
    resp = client.get(reverse('admin_settings'))
    check("Admin: Configuracion Institucional GET 200", resp.status_code == 200)

    # Reportes globales de resultados
    resp = client.get(reverse('admin_reportes'))
    check("Admin: Reportes y Metricas Globales GET 200", resp.status_code == 200)

    # ------------------------------------------------------------------
    # 4. MODULO DOCENTE (M02, M07, M08, M09, M10, M12)
    # ------------------------------------------------------------------
    print("\n4. VERIFICANDO MODULO DOCENTE...")
    s = client.session
    s['demo_profile'] = {
        'id': docente_user.id,
        'role': 'docente',
        'name': docente_user.nombre_completo,
        'code': docente_user.codigo_institucional,
        'email': docente_user.email,
        'redirect': 'docente_dashboard'
    }
    s.save()

    # Dashboard Docente
    resp = client.get(reverse('docente_dashboard'))
    check("Docente: Dashboard Academico GET 200", resp.status_code == 200)

    # Areas y Competencias
    resp = client.get(reverse('docente_competencias'))
    check("Docente: Modulo Competencias GET 200", resp.status_code == 200)

    # Banco de Items
    resp = client.get(reverse('docente_preguntas'))
    check("Docente: Banco de Preguntas GET 200", resp.status_code == 200)

    # Crear nueva pregunta manual
    categoria_sample = Categoria.objects.first()
    resp = client.post(reverse('docente_pregunta_nueva'), {
        'categoria_id': categoria_sample.id,
        'tipo': 'MULTIPLE',
        'enunciado': 'Pregunta de prueba generada en auditoria?',
        'texto_base': 'Contexto de prueba',
        'puntaje_base': '1.00',
        'retroalimentacion': 'Explicacion de prueba',
        'opcion_A_texto': 'Alternativa 1',
        'opcion_B_texto': 'Alternativa 2',
        'opcion_C_texto': 'Alternativa 3',
        'opcion_D_texto': 'Alternativa 4',
        'correcta_letra': 'A'
    }, follow=True)
    check("Docente: Crear reactivo pedagogico manual", resp.status_code == 200)

    # Descarga de Plantilla Oficial Excel
    resp = client.get(reverse('descargar_plantilla_excel'))
    check("Docente: Descarga Plantilla Excel (.xlsx) 200", resp.status_code == 200)
    check("Docente: Header Content-Type Excel (.xlsx)", resp.get('Content-Type') == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # Validar estructura del archivo Excel generado
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    check("Docente: Plantilla Excel incluye pestana 'Banco de Preguntas'", 'Banco de Preguntas' in wb.sheetnames)
    check("Docente: Plantilla Excel incluye pestana 'Instrucciones para Docentes'", 'Instrucciones para Docentes' in wb.sheetnames)

    # Asistente de importacion (Fase 1 y 2) con el Excel generado
    excel_file = SimpleUploadedFile(
        'plantilla_prueba.xlsx',
        resp.content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp = client.post(reverse('docente_importaciones'), {
        'accion': 'validar',
        'archivo_excel': excel_file
    }, follow=True)
    check("Docente: Importador valida archivo Excel", resp.status_code == 200)
    staging = client.session.get('staging_importacion')
    check("Docente: Importador extrae reactivos validos", staging is not None and len(staging.get('filas_validas', [])) == 4)

    # Programador de Cuestionarios
    resp = client.get(reverse('docente_cuestionarios'))
    check("Docente: Lista de Cuestionarios GET 200", resp.status_code == 200)

    cuest_sample = Cuestionario.objects.first()
    resp = client.get(reverse('docente_cuestionario_preguntas', kwargs={'pk': cuest_sample.id}))
    check("Docente: Gestion de Ponderaciones y Preguntas GET 200", resp.status_code == 200)

    resp = client.get(reverse('docente_intentos'))
    check("Docente: Metricas y Resultados por Cuestionario GET 200", resp.status_code == 200)

    # ------------------------------------------------------------------
    # 5. MODULO ESTUDIANTE (M02, M11, M12)
    # ------------------------------------------------------------------
    print("\n5. VERIFICANDO MODULO ESTUDIANTE...")
    s = client.session
    s['demo_profile'] = {
        'id': estudiante_user.id,
        'role': 'estudiante',
        'name': estudiante_user.nombre_completo,
        'code': estudiante_user.codigo_institucional,
        'email': estudiante_user.email,
        'redirect': 'estudiante_catalogo'
    }
    s.save()

    # Catalogo de Simulacros
    resp = client.get(reverse('estudiante_catalogo'))
    check("Estudiante: Catalogo de Simulacros GET 200", resp.status_code == 200)

    # Sala de Espera / Ficha Tecnica
    resp = client.get(reverse('estudiante_sala_espera', kwargs={'pk': cuest_sample.id}))
    check("Estudiante: Sala de Espera GET 200", resp.status_code == 200)
    content_espera = resp.content.decode('utf-8')
    check("Estudiante: Identidad institucional bloqueada", 'Bloqueado' in content_espera and estudiante_user.nombre_completo in content_espera)

    # Iniciar Intento Oficial
    resp = client.post(reverse('estudiante_iniciar', kwargs={'pk': cuest_sample.id}), follow=True)
    check("Estudiante: Iniciar Intento redirige al simulacro", resp.status_code == 200)
    intento_activo = Intento.objects.filter(cuestionario=cuest_sample, codigo_estudiante=estudiante_user.codigo_institucional).order_by('-id').first()
    check("Estudiante: Intento creado en BD", intento_activo is not None and intento_activo.estado == 'En Progreso')

    # Vista del Simulacro en Vivo
    resp = client.get(reverse('estudiante_simulacro', kwargs={'intento_id': intento_activo.id}))
    check("Estudiante: Simulacro en Vivo GET 200", resp.status_code == 200)
    content_sim = resp.content.decode('utf-8')
    check("Estudiante: Script de motor simulacro.js incluido con versionado", 'simulacro.js?v=' in content_sim)
    check("Estudiante: Botones de navegacion siguiente/anterior presentes", 'next-question-btn' in content_sim and 'prev-question-btn' in content_sim)

    # Liquidacion y Entrega del Simulacro
    # Contestar 2 preguntas y dejar las demas vacias
    cp_list = cuest_sample.cuestionario_preguntas.all().select_related('pregunta')
    post_answers = {}
    for i, cp in enumerate(cp_list):
        if i < 2:
            opc = cp.pregunta.opciones.filter(es_correcta=True).first()
            if opc:
                post_answers[f'pregunta_{cp.pregunta.id}'] = opc.letra

    resp = client.post(reverse('estudiante_entregar', kwargs={'intento_id': intento_activo.id}), post_answers, follow=True)
    check("Estudiante: Entregar Simulacro redirige al informe 200", resp.status_code == 200)
    intento_activo.refresh_from_db()
    check("Estudiante: Intento marcado como 'Finalizado'", intento_activo.estado == 'Finalizado')
    check("Estudiante: Calificacion real calculada en BD", intento_activo.calificacion_final is not None and intento_activo.calificacion_final > Decimal('0.00'))

    # Informe Evaluativo y Retroalimentacion
    resp = client.get(reverse('estudiante_resultado', kwargs={'intento_id': intento_activo.id}))
    check("Estudiante: Informe de Resultados GET 200", resp.status_code == 200)
    content_res = resp.content.decode('utf-8')
    check("Estudiante: Desglose de competencias y preguntas en el informe", 'Retroalimentaci' in content_res)

    # ------------------------------------------------------------------
    # 6. MODELO DE DATOS Y ARQUITECTURA
    # ------------------------------------------------------------------
    print("\n6. VERIFICANDO ARQUITECTURA Y MODELO DE DATOS...")
    resp = client.get(reverse('modelo_datos'))
    check("Arquitectura: Vista Modelo de Datos GET 200", resp.status_code == 200)

    # ------------------------------------------------------------------
    # RESUMEN GENERAL
    # ------------------------------------------------------------------
    print("\n======================================================================")
    print(f"RESUMEN FINAL: {len(successes)} pruebas exitosas, {len(errors)} errores encontrados.")
    print("======================================================================")

    if errors:
        print("\nERRORES DETECTADOS:")
        for err in errors:
            print(f"  [ERROR] {err}")
        return False
    else:
        print("\n[OK] TODOS LOS MODULOS FUNCIONALES ESTAN OPERATIVOS Y SIN ERRORES (100% OK).")
        return True

if __name__ == '__main__':
    audit_system()
